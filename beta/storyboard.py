"""Reads a storyboard (SB) into plain text for content mapping.

Storyboards are usually Word documents with one table per screen (screen no, OST, VO/audio
script, developer notes); PDFs, PowerPoint and Excel storyboards are read too. Table rows
are kept as "cell | cell | cell" lines so screen numbers stay next to their content.
"""
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

SUPPORTED = {".docx", ".doc", ".rtf", ".odt", ".pdf", ".pptx", ".ppt", ".xlsx", ".xls"}


def _clean(t):
    return re.sub(r"[ \t]+", " ", (t or "").replace(" ", " ")).strip()


def _docx(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(str(path))
    out = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = _clean(Paragraph(block, doc).text)
            if t:
                out.append(t)
        elif tag == "tbl":
            out.append("")
            for row in Table(block, doc).rows:
                cells, seen = [], set()
                for c in row.cells:
                    if id(c._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(c._tc))
                    cells.append(_clean(c.text).replace("\n", " / "))
                if any(cells):
                    out.append(" | ".join(cells))
            out.append("")
    return "\n".join(out)


def _pdf(path):
    p = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, timeout=300)
    pages = p.stdout.split("\f")
    return "\n".join(f"=== SB page {i} ===\n{t}" for i, t in enumerate(pages, 1) if t.strip())


def _pptx(path):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        out.append(f"=== SB slide {i} ===")
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = _clean(sh.text_frame.text)
                if t:
                    out.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append(" | ".join(_clean(c.text) for c in row.cells))
        if slide.has_notes_slide:
            n = _clean(slide.notes_slide.notes_text_frame.text)
            if n:
                out.append("Notes: " + n)
    return "\n".join(out)


def _xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Sheet {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [_clean(str(v)) for v in row if v is not None and str(v).strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


NOTES_PDF = ('pdf:impress_pdf_Export:{"ExportNotesPages":{"type":"boolean","value":"true"},'
             '"ExportOnlyNotesPages":{"type":"boolean","value":"true"}}')


def _convert(path, fmt, target=None):
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", target or fmt, "--outdir", str(tmp), str(path)],
                   capture_output=True, timeout=300)
    hits = list(tmp.glob("*." + fmt))
    if not hits:
        raise ValueError(f"Could not convert the storyboard ({path.suffix}) to {fmt}")
    return hits[0]


def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Storyboard type {ext} is not supported. Use Word, PDF, PowerPoint or Excel.")
    if ext in (".doc", ".rtf", ".odt"):
        return _docx(_convert(path, "docx"))
    if ext == ".ppt":
        return _pptx(_convert(path, "pptx"))
    if ext == ".xls":
        return _xlsx(_convert(path, "xlsx"))
    return {".docx": _docx, ".pdf": _pdf, ".pptx": _pptx, ".xlsx": _xlsx}[ext](path)


# ---------------------------------------------------------------- page pictures (for "Not as per SB" evidence)
DPI = 150
WORD_RE = re.compile(r"[\w'’‘]+")


def _norm(w):
    """Words compare without case and apostrophes (You’ve = You've = youve)."""
    return re.sub(r"[^\w]+", "", w.lower())


def _render_pdf(pdf, out_dir: Path, prefix, max_pages):
    """Pictures (out_dir/<prefix>N.png) and word positions {"N": [[word, x, y, w, h], …]} of one PDF."""
    import html
    tmp = out_dir / f"_{prefix}"
    tmp.mkdir(parents=True, exist_ok=True)
    subprocess.run(["pdftoppm", "-r", str(DPI), "-png", "-l", str(max_pages), str(pdf), str(tmp / "p")],
                   capture_output=True, timeout=600)
    for p in tmp.glob("p-*.png"):   # pdftoppm pads the number (p-01.png): store as <prefix>1.png
        n = int(p.stem.split("-")[-1])
        p.replace(out_dir / f"{prefix}{n}.png")
    shutil.rmtree(tmp, ignore_errors=True)
    bbox = subprocess.run(["pdftotext", "-bbox-layout", "-l", str(max_pages), str(pdf), "-"],
                          capture_output=True, text=True, timeout=600).stdout
    k = DPI / 72.0
    pages = {}
    for n, body in enumerate(re.findall(r"<page[^>]*>(.*?)</page>", bbox, re.S), 1):
        words = []
        for m in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>', body):
            x0, y0, x1, y1 = (float(m.group(i)) * k for i in range(1, 5))
            words.append([html.unescape(m.group(5)), round(x0, 1), round(y0, 1), round(x1 - x0, 1), round(y1 - y0, 1)])
        pages[str(n)] = words
    return pages


def render_pages(path: Path, out_dir: Path, max_pages=300):
    """Saves every SB page as a picture (out_dir/pageN.png) and the position of each word on it
    (out_dir/words.json). A PowerPoint storyboard also gets its notes pages – the slide with its speaker notes
    (the VO) underneath – as out_dir/notesN.png and out_dir/words_notes.json, so issues in the VO can be shown
    and marked too. Returns the page count."""
    import json
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf = path if path.suffix.lower() == ".pdf" else _convert(path, "pdf")
    pages = _render_pdf(pdf, out_dir, "page", max_pages)
    (out_dir / "words.json").write_text(json.dumps(pages))
    if path.suffix.lower() in (".pptx", ".ppt", ".odp"):
        try:
            notes_pdf = _convert(path, "pdf", NOTES_PDF)
            notes = _render_pdf(notes_pdf, out_dir, "notes", max_pages)
            if len(notes) == len(pages):   # one notes page per slide picture: same numbering
                (out_dir / "words_notes.json").write_text(json.dumps(notes))
        except Exception:
            pass
    return len(pages)


def _split(words):
    """PDF words split into the same pieces the quote is split into, each with its share of the box:
    "authorization(P" → "authorization" + "P" (a missing space must never stop a quote from being found)."""
    out = []
    for w in words:
        t = w[0]
        parts = list(WORD_RE.finditer(t))
        if len(parts) <= 1:
            out.append(w)
            continue
        n = max(1, len(t))
        for m in parts:
            out.append([m.group(0), w[1] + w[3] * m.start() / n, w[2], w[3] * (m.end() - m.start()) / n, w[4]])
    return out


def _match(words, text):
    """(selected words, score 0..1) of the best place on one page where `text` is. Word by word first; if that
    fails, letter by letter ignoring spaces and punctuation, so "authorization(P A)" is found as
    "authorization(PA)" or "authorization (P A)" too."""
    want = [_norm(w) for w in WORD_RE.findall(text or "") if _norm(w)][:40]
    if not want or not words:
        return [], 0.0
    words = _split(words)
    norm = [_norm(w[0]) for w in words]
    n = len(want)
    best, best_score = [], 0.0
    for i in range(len(norm)):
        if norm[i] != want[0] and (n < 2 or norm[i] != want[1]):
            continue
        j, hit, sel = i, 0, []
        for wv in want:
            for jj in range(j, min(j + 4, len(norm))):
                if norm[jj] == wv:
                    hit += 1
                    sel.append(words[jj])
                    j = jj + 1
                    break
        if hit / n > best_score:
            best, best_score = sel, hit / n
    if best_score >= 0.8:
        return best, best_score
    # letters only: the whole page as one string, each letter knowing its word
    target = "".join(want)
    if len(target) < 3:
        return best, best_score
    chars, owner = [], []
    for k, t in enumerate(norm):
        chars.append(t)
        owner += [k] * len(t)
    page = "".join(chars)
    at = page.find(target)
    if at < 0 and len(target) > 12:   # most of a long quote is enough
        for cut in (len(target) * 3 // 4, len(target) // 2):
            at = page.find(target[:cut])
            if at >= 0:
                target = target[:cut]
                break
    if at >= 0:
        ks = sorted(set(owner[at:at + len(target)]))
        return [words[k] for k in ks], max(best_score, 0.95 if len(target) == len("".join(want)) else 0.7)
    return best, best_score


def _box(sel):
    x0 = min(w[1] for w in sel); y0 = min(w[2] for w in sel)
    x1 = max(w[1] + w[3] for w in sel); y1 = max(w[2] + w[4] for w in sel)
    return [x0, y0, x1 - x0, y1 - y0]


def find_on_pages(words_by_page, text, page_hint=None):
    """The SB page and box that best match `text` (the SB wording quoted in an issue).
    Returns (page number, [x, y, w, h] or None, score 0..1). The page the issue names wins whenever the words are there."""
    if not [w for w in WORD_RE.findall(text or "") if _norm(w)]:
        return (int(page_hint) if page_hint else None), None, 0.0
    if page_hint and str(page_hint) in words_by_page:
        sel, sc = _match(words_by_page[str(page_hint)], text)
        if sel and sc >= 0.6:
            return int(page_hint), _box(sel), sc
    best = (None, None, 0.0)
    for pg, words in words_by_page.items():
        sel, sc = _match(words, text)
        if sel and sc > best[2]:
            best = (int(pg), _box(sel), sc)
    if best[2] >= 0.6:
        return best
    return (int(page_hint) if page_hint else best[0]), None, best[2]


def quote_lines(words, text):
    """Boxes of the quoted words on one page, one box per printed line (a quote that wraps is marked line by line)."""
    sel, sc = _match(words, text)
    if not sel or sc < 0.6:
        return []
    lines = []
    for w in sel:   # group by line: same top (within half a line height)
        for ln in lines:
            if abs(ln[0][2] - w[2]) < max(4, w[4] * 0.5):
                ln.append(w)
                break
        else:
            lines.append([w])
    return [_box(ln) for ln in lines]

