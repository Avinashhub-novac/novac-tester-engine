"""Excel export of a finished run."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILL = {"error": "F8D7DA", "warning": "FFF3CD", "info": "E2E3E5",
        "fail": "F8D7DA", "warn": "FFF3CD", "pass": "D1E7DD"}
HEAD = PatternFill("solid", fgColor="1F2937")


def _sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build(meta, results, base_url=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Novac Tester report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["File", meta.get("name", "")])
    ws.append(["Type", "SCORM (Storyline)" if meta.get("kind") == "scorm" else "Video"])
    ws.append(["Uploaded by", meta.get("uploader", "")])
    ws.append(["Course title", results.get("course_title", "")])
    ws.append([])
    start = ws.max_row + 1
    _rows = [[v["name"], v["status"].upper(), v["errors"], v["warnings"], v.get("skipped_reason", "")]
             for v in results["summary"].values()]
    ws.append(["Check", "Result", "Errors", "Warnings", "Note"])
    for c in ws[start]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
    for r in _rows:
        ws.append(r)
        cell = ws.cell(row=ws.max_row, column=2)
        if r[1].lower() in FILL:
            cell.fill = PatternFill("solid", fgColor=FILL[r[1].lower()])
    for col, w in zip("ABCDE", (22, 12, 10, 10, 60)):
        ws.column_dimensions[col].width = w

    names = {k: v["name"] for k, v in results["summary"].items()}
    order = {"error": 0, "warning": 1, "info": 2}
    rows = []
    for f in sorted(results["findings"], key=lambda f: (order[f["severity"]], f["check"])):
        rows.append([f["severity"].upper(), names.get(f["check"], f["check"]), f.get("slide_label") or f.get("time", ""),
                     f["message"], f.get("detail", ""), f.get("suggestion", ""),
                     (base_url + "/" + f["screenshot"]) if f.get("screenshot") and base_url else f.get("screenshot", "")])
    ws2 = wb.create_sheet("Findings")
    _sheet(ws2, ["Severity", "Check", "Slide / time", "Problem", "Detail", "Suggestion", "Screenshot"], rows,
           (10, 14, 28, 50, 60, 30, 40))
    for row in ws2.iter_rows(min_row=2):
        sev = (row[0].value or "").lower()
        if sev in FILL:
            row[0].fill = PatternFill("solid", fgColor=FILL[sev])

    if results.get("slides"):
        per = {}
        for f in results["findings"]:
            per.setdefault(f.get("slide_id"), []).append(f)
        srows = []
        for s in results["slides"]:
            fs = per.get(s["id"], [])
            srows.append([s["label"], "yes" if s["reached"] else "no", len(s.get("audio", [])), s.get("duration", ""),
                          sum(f["severity"] == "error" for f in fs), sum(f["severity"] == "warning" for f in fs),
                          "; ".join(sorted({names.get(f["check"], f["check"]) for f in fs if f["severity"] != "info"}))])
        _sheet(wb.create_sheet("Slides"), ["Slide", "Reached by Next", "Audio files", "Timeline (s)", "Errors",
                                            "Warnings", "Failed checks"], srows, (40, 14, 11, 12, 8, 9, 50))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
